"""Paso 4 de la skill remision-produccion-tabs: antes de dar la orden por buena se vuelve a leer el Excel del listado y se compara con el PDF.
Cuenta una unidad por cada fila con TALLA (columna D, filas 8-507 de cada pestaña) y la compara con el «Total unidades» del PDF
(o, si el PDF no lo trae, con la suma de las cantidades de sus ítems)."""
from __future__ import annotations

import re
from pathlib import Path

PRIMERA, ULTIMA = 8, 507
HOJAS_FIJAS = {'REFERENCIA', 'BASE_DATOS'}


def total_del_pdf(texto: str) -> int | None:
    """«Total unidades: 245» (con puntos de miles o sin ellos). Devuelve None si el PDF no lo dice."""
    m = re.search(r'total\s+unidades\s*[:\-]?\s*(\d[\d.,]*)', texto or '', re.I)
    if not m:
        return None
    digitos = re.sub(r'\D', '', m.group(1))
    return int(digitos) if digitos else None


def unidades_por_hoja(excel: Path) -> dict[str, int]:
    import io
    from openpyxl import load_workbook
    libro = load_workbook(io.BytesIO(Path(excel).read_bytes()), read_only=True, data_only=True)
    try:
        cuenta = {}
        for hoja in libro.worksheets:
            if hoja.title.upper() in HOJAS_FIJAS:
                continue
            n = 0
            for fila in hoja.iter_rows(min_row=PRIMERA, max_row=ULTIMA, min_col=4, max_col=4, values_only=True):
                if fila and fila[0] not in (None, ''):
                    n += 1
            cuenta[hoja.title] = n
        return cuenta
    finally:
        libro.close()


def verificar(pdf_texto: str, items: list[dict], excel: Path) -> dict:
    """{'ok', 'pdf', 'excel', 'hojas', 'mensaje'}. `items` son los ítems que leyó el PDF (con 'cantidad')."""
    por_hoja = unidades_por_hoja(excel)
    en_excel = sum(por_hoja.values())
    de_items = sum(int(it.get('cantidad') or 0) for it in items or [])
    del_pdf = total_del_pdf(pdf_texto)
    esperado = del_pdf if del_pdf is not None else de_items
    origen = 'el «Total unidades» del PDF' if del_pdf is not None else 'la suma de los ítems del PDF'
    problemas = []
    if en_excel != esperado:
        problemas.append(f'el Excel tiene {en_excel} unidades y {origen} dice {esperado}')
    if del_pdf is not None and de_items and de_items != del_pdf:
        problemas.append(f'los ítems leídos suman {de_items} y el PDF dice {del_pdf} (puede faltar algún ítem)')
    if problemas:
        return {'ok': False, 'pdf': esperado, 'excel': en_excel, 'hojas': por_hoja, 'mensaje': 'NO CUADRA: ' + '; '.join(problemas)}
    return {'ok': True, 'pdf': esperado, 'excel': en_excel, 'hojas': por_hoja,
            'mensaje': f'Verificado: {en_excel} unidades en {len(por_hoja)} pestaña(s), igual que {origen}'}
