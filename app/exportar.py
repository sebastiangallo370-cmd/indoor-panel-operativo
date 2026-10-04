"""Descargas en Excel (o CSV) de inventario, movimientos, cartera y producción."""
import csv
import io
import re
from datetime import datetime, timedelta, timezone
from typing import Any, Callable

from fastapi import APIRouter, HTTPException, Request, Response

from app import inventario_api, permisos

router = APIRouter(prefix='/api/exportar', tags=['exportar'])
_leer_produccion: Callable[[], dict] = lambda: {}
_autenticar: Callable = lambda request: None

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
MODULO_DE = {'inventario': 'inventario', 'movimientos': 'inventario', 'cartera': 'cartera', 'produccion': 'produccion'}


def configurar(leer_produccion: Callable[[], dict], autenticar: Callable) -> None:
    global _leer_produccion, _autenticar
    _leer_produccion, _autenticar = leer_produccion, autenticar


def _numero(valor: Any):
    try:
        return float(valor) if valor not in ('', None) else None
    except (TypeError, ValueError):
        return None


def _hoja_inventario() -> list[tuple[str, list[str], list[list]]]:
    items = inventario_api._with_sublimacion(inventario_api._inventory_payload()).get('items') or []
    por_categoria: dict[str, list[dict]] = {}
    for item in items:
        por_categoria.setdefault(item.get('categoria_label') or item.get('categoria') or 'Sin categoría', []).append(item)
    hojas = []
    for categoria, lista in por_categoria.items():
        campos = []
        for item in lista:
            for clave in (item.get('campos') or {}):
                if clave not in campos:
                    campos.append(clave)
        es_tela = any(i.get('categoria') in ('BODEGA TELA', 'RETAL CANASTAS') for i in lista)
        cabecera = ['Nombre', 'Total (MTS/unid.)'] + (['Rollos', 'Rollos empezados'] if es_tela else []) + campos
        filas = []
        for item in lista:
            fila = [item.get('nombre', ''), _numero(item.get('total'))]
            if es_tela:
                estados = item.get('roll_statuses') or []
                fila += [len(item.get('roll_values') or []) or None, sum(1 for s in estados if s == 'started') or None]
            fila += [(item.get('campos') or {}).get(c, '') for c in campos]
            filas.append(fila)
        hojas.append((categoria, cabecera, filas))
    return hojas


def _hoja_movimientos() -> list[tuple[str, list[str], list[list]]]:
    zona = timezone(timedelta(hours=-5))
    filas = []
    for m in reversed(inventario_api._load_movements()):
        try:
            fecha = datetime.fromisoformat(m.get('fecha', '')).astimezone(zona).strftime('%Y-%m-%d %H:%M')
        except ValueError:
            fecha = m.get('fecha', '')
        filas.append([fecha, m.get('type', ''), m.get('name', ''), m.get('code', ''), _numero(m.get('mts')),
                      m.get('rolls'), m.get('bodega', ''), m.get('source', '')])
    return [('Movimientos', ['Fecha (Bogotá)', 'Tipo', 'Tela', 'Código', 'MTS', 'Rollos', 'Bodega', 'Origen'], filas)]


def _hoja_cartera() -> list[tuple[str, list[str], list[list]]]:
    from app import cartera_api
    datos = cartera_api._load()
    pagado: dict[str, float] = {}
    for c in datos.get('comprobantes') or []:
        if not c.get('anulado'):
            pagado[str(c.get('cotizacionNumero'))] = pagado.get(str(c.get('cotizacionNumero')), 0) + (_numero(c.get('valor')) or 0)
    docs = []
    for d in datos.get('documentos') or []:
        total = _numero(d.get('total')) or 0
        abonado = (_numero(d.get('pagadoImportado')) or 0) + pagado.get(str(d.get('numero')), 0)
        docs.append([d.get('numero'), d.get('cliente'), d.get('vendedor'), d.get('fechaCreacion'), d.get('fechaEntrega'),
                     d.get('formaPago'), d.get('estado'), total, abonado, total - abonado, d.get('factura', '')])
    pagos = [[c.get('fecha'), c.get('cotizacionNumero'), c.get('cliente'), _numero(c.get('valor')), c.get('medio'),
              c.get('referencia'), c.get('recibio'), 'Anulado' if c.get('anulado') else ''] for c in datos.get('comprobantes') or []]
    return [
        ('Cartera', ['Cotización', 'Cliente', 'Vendedor', 'Creación', 'Entrega', 'Forma de pago', 'Estado', 'Total', 'Pagado', 'Saldo', 'Factura'], docs),
        ('Pagos', ['Fecha', 'Cotización', 'Cliente', 'Valor', 'Medio', 'Referencia', 'Recibió', 'Nota'], pagos),
    ]


def _hoja_produccion() -> list[tuple[str, list[str], list[list]]]:
    datos = _leer_produccion()
    cabecera = [str(h or '').strip() or f'Col {i + 1}' for i, h in enumerate(datos.get('headers') or [])]
    filas = [[(fila.get('values') or [])[i] if i < len(fila.get('values') or []) else '' for i in range(len(cabecera))]
             for fila in datos.get('rows') or []]
    return [('Producción', cabecera, filas)]


FUENTES = {'inventario': _hoja_inventario, 'movimientos': _hoja_movimientos, 'cartera': _hoja_cartera, 'produccion': _hoja_produccion}


def _celda_segura(valor):
    # Evita que Excel interprete texto del Sheet como fórmula (=, +, -, @).
    if isinstance(valor, str) and valor[:1] in ('=', '+', '-', '@') and not re.fullmatch(r'-?[\d.,]+', valor):
        return "'" + valor
    return valor


def _xlsx(hojas) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    libro = Workbook()
    libro.remove(libro.active)
    usados = set()
    for nombre, cabecera, filas in hojas:
        titulo = re.sub(r'[\[\]\*\?/\\:]', ' ', nombre)[:31] or 'Hoja'
        while titulo in usados:
            titulo = titulo[:28] + '_' + str(len(usados))
        usados.add(titulo)
        hoja = libro.create_sheet(titulo)
        hoja.append(cabecera)
        for fila in filas:
            hoja.append([_celda_segura(v) for v in fila])
        for celda in hoja[1]:
            celda.font = Font(bold=True, color='FFFFFF')
            celda.fill = PatternFill('solid', fgColor='2B3A27')
            celda.alignment = Alignment(vertical='center', wrap_text=True)
        hoja.freeze_panes = 'A2'
        if filas:
            hoja.auto_filter.ref = hoja.dimensions
        for i, titulo_col in enumerate(cabecera, 1):
            largo = max([len(str(titulo_col))] + [len(str(f[i - 1])) for f in filas[:300] if i - 1 < len(f) and f[i - 1] is not None])
            hoja.column_dimensions[get_column_letter(i)].width = min(max(largo + 2, 9), 48)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _csv(hojas) -> bytes:
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=';')
    for indice, (nombre, cabecera, filas) in enumerate(hojas):
        if indice:
            escritor.writerow([])
        if len(hojas) > 1:
            escritor.writerow([nombre])
        escritor.writerow(cabecera)
        for fila in filas:
            escritor.writerow([_celda_segura(v) if v is not None else '' for v in fila])
    return ('﻿' + salida.getvalue()).encode('utf-8')  # BOM: Excel abre bien las tildes


@router.get('/{tipo}')
def exportar(tipo: str, request: Request, formato: str = 'xlsx'):
    usuario = _autenticar(request)
    if tipo not in FUENTES:
        raise HTTPException(404, 'Esa exportación no existe')
    if formato not in ('xlsx', 'csv'):
        raise HTTPException(400, 'El formato debe ser xlsx o csv')
    if not permisos.puede(usuario, MODULO_DE[tipo], 'exportar'):
        raise HTTPException(403, 'Tu perfil no tiene permiso para exportar esta información')
    hojas = FUENTES[tipo]()
    hoy = datetime.now(timezone(timedelta(hours=-5))).strftime('%Y-%m-%d')
    nombre = f'{tipo}_{hoy}.{formato}'
    cuerpo, tipo_mime = (_xlsx(hojas), XLSX) if formato == 'xlsx' else (_csv(hojas), 'text/csv; charset=utf-8')
    return Response(cuerpo, media_type=tipo_mime, headers={'Content-Disposition': f'attachment; filename="{nombre}"', 'Cache-Control': 'no-store'})
