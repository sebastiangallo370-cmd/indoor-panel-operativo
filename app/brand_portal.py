"""Opt-in client workspace. No production writes and no public orders by default."""
import hashlib
import io
import json
import os
import secrets
import time
import unicodedata
import zipfile
from contextlib import contextmanager, closing
from pathlib import Path
from urllib.parse import urlsplit

from fastapi import Body, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, Response
from openpyxl import load_workbook


def normalized(value):
    return ''.join(c for c in unicodedata.normalize('NFD', str(value or '')) if unicodedata.category(c) != 'Mn').strip().upper()


def validate_players(players):
    if not isinstance(players, list) or len(players) > 500:
        raise HTTPException(422, 'Máximo 500 jugadores por listado')
    result, seen = [], set()
    for i, item in enumerate(players, 1):
        if not isinstance(item, dict):
            raise HTTPException(422, f'Jugador {i}: formato inválido')
        name, number, size = (str(item.get(k, '')).strip() for k in ('name', 'number', 'size'))
        if not name or len(name) > 80 or not number.isdigit() or len(number) > 3 or not size or len(size) > 12:
            raise HTTPException(422, f'Jugador {i}: completa nombre, dorsal (0–999) y talla')
        if int(number) in seen:
            raise HTTPException(422, f'Dorsal repetido: {number}')
        seen.add(int(number))
        result.append(dict(name=name, number=number, size=size.upper()))
    return result


def register_portal(app, connect, authenticate, read_production, row_files, excel_designs):
    directory = Path(__file__).parent

    @app.middleware('http')
    async def protect_portal_writes(request: Request, call_next):
        if request.url.path.startswith(('/api/equipo/', '/api/estudio/')):
            if request.method not in ('GET', 'HEAD', 'OPTIONS'):
                origin = request.headers.get('origin')
                if request.headers.get('sec-fetch-site') == 'cross-site' or (origin and urlsplit(origin).netloc != request.headers.get('host')):
                    from fastapi.responses import JSONResponse
                    return JSONResponse({'detail': 'Origen no permitido'}, status_code=403)
                if int(request.headers.get('content-length', '0') or 0) > 2_100_000:
                    from fastapi.responses import JSONResponse
                    return JSONResponse({'detail': 'Solicitud demasiado grande'}, status_code=413)
        response = await call_next(request)
        if request.url.path.startswith(('/api/equipo/', '/api/estudio')):
            response.headers['Cache-Control'] = 'private, no-store'
            response.headers['Referrer-Policy'] = 'no-referrer'
        return response

    @contextmanager
    def database():
        db = connect()
        db.executescript('''
          CREATE TABLE IF NOT EXISTS team_workspaces (
            source_row INTEGER PRIMARY KEY, identity TEXT NOT NULL, version INTEGER NOT NULL DEFAULT 0,
            designs TEXT NOT NULL DEFAULT '[]', decision TEXT NOT NULL DEFAULT 'pending',
            players TEXT NOT NULL DEFAULT '[]', revision INTEGER NOT NULL DEFAULT 0,
            published INTEGER NOT NULL DEFAULT 0, sport TEXT NOT NULL DEFAULT 'Fútbol');
          CREATE TABLE IF NOT EXISTS team_links (
            digest TEXT PRIMARY KEY, source_row INTEGER NOT NULL, expires REAL NOT NULL);
          CREATE TABLE IF NOT EXISTS team_images (digest TEXT PRIMARY KEY, mime TEXT NOT NULL, content BLOB NOT NULL);
          CREATE TABLE IF NOT EXISTS team_events (
            id INTEGER PRIMARY KEY, source_row INTEGER NOT NULL, action TEXT NOT NULL,
            actor TEXT NOT NULL, detail TEXT NOT NULL, created REAL NOT NULL);
          CREATE TABLE IF NOT EXISTS team_requests (
            id INTEGER PRIMARY KEY, source_row INTEGER NOT NULL, kind TEXT NOT NULL,
            detail TEXT NOT NULL, created REAL NOT NULL, status TEXT NOT NULL DEFAULT 'Pendiente');
        ''')
        try:
            with db:
                yield db
        finally:
            db.close()

    def admin(user=Depends(authenticate)):
        with closing(connect()) as db:
            profile = db.execute('SELECT process FROM users WHERE name=? COLLATE NOCASE', (user,)).fetchone()
        role = normalized(profile['process'] if profile else 'Administración' if user == os.getenv('APP_USER', 'indoor') else '')
        if role not in ('ADMINISTRACION', 'COMERCIAL', 'COMERCIALES', 'ASISTENTE COMERCIAL', 'ASISTENTES COMERCIALES'):
            raise HTTPException(403, 'Solo Administración y asistentes comerciales')
        return user

    def row_info(source_row):
        data = read_production()
        row = next((r for r in data['rows'] if r['source_row'] == source_row), None)
        if row is None:
            raise HTTPException(404, 'Pedido no encontrado')
        fields = {normalized(h): str(v or '').strip() for h, v in zip(data['headers'], row['values'])}
        result = dict(row=source_row, order=fields.get('ORDEN', ''), client=fields.get('NOMBRE DEL CLIENTE', ''),
                      project=fields.get('NOMBRE PROYECTO', ''), reference=fields.get('REFERENCIA', ''),
                      quantity=fields.get('CANTIDAD', ''), due=fields.get('FECHA DE ENTREGA', ''),
                      process=row.get('current_process', {'label': 'Sin iniciar', 'state': 'pending'}))
        if not result['order'] or not result['client']:
            raise HTTPException(422, 'La fila no corresponde a un pedido programado')
        return result

    def identity(info):
        return json.dumps([info[k] for k in ('order', 'client', 'reference')], ensure_ascii=False)

    def workspace(db, source_row, info):
        record = db.execute('SELECT * FROM team_workspaces WHERE source_row=?', (source_row,)).fetchone()
        if record and record['identity'] != identity(info):
            raise HTTPException(409, 'Cambió la identidad de esta fila. Revoca los enlaces y revisa el pedido antes de compartirlo.')
        if not record:
            db.execute('INSERT INTO team_workspaces(source_row,identity) VALUES(?,?)', (source_row, identity(info)))
            record = db.execute('SELECT * FROM team_workspaces WHERE source_row=?', (source_row,)).fetchone()
        return record

    def event(db, source_row, action, actor, detail=''):
        db.execute('INSERT INTO team_events(source_row,action,actor,detail,created) VALUES(?,?,?,?,?)',
                   (source_row, action, actor, detail, time.time()))

    def token_row(token):
        with database() as db:
            link = db.execute('SELECT source_row FROM team_links WHERE digest=? AND expires>?',
                              (hashlib.sha256(token.encode()).hexdigest(), time.time())).fetchone()
        if not link:
            raise HTTPException(404, 'Enlace vencido o revocado. Solicita uno nuevo a tu asesor.')
        return link['source_row']

    def details(source_row, token=None):
        info = row_info(source_row)
        with database() as db:
            record = workspace(db, source_row, info)
            info.update(version=record['version'], decision=record['decision'], players=json.loads(record['players']),
                        revision=record['revision'], published=bool(record['published']), sport=record['sport'])
            prefix = f'/api/equipo/{token}' if token else f'/api/estudio/{source_row}'
            info['designs'] = [dict(number=d['number'], url=f"{prefix}/imagen/{d['digest']}") for d in json.loads(record['designs'])]
            info['history'] = [dict(r) for r in db.execute('SELECT action,actor,detail,created FROM team_events WHERE source_row=? ORDER BY id DESC LIMIT 30', (source_row,))]
        return info

    def page(mode):
        html = (directory / 'portal.html').read_text(encoding='utf-8').replace('__MODE__', mode)
        return HTMLResponse(html, headers={'Referrer-Policy': 'no-referrer', 'Cache-Control': 'no-store', 'X-Robots-Tag': 'noindex'})

    @app.get('/estudio')
    def studio_page(user=Depends(admin)):
        return page('admin')

    @app.get('/catalogo')
    def catalog_page():
        return page('catalog')

    @app.get('/equipo/{token}')
    def team_page(token: str):
        token_row(token)
        return page('client')

    @app.get('/portal.css')
    def css():
        return FileResponse(directory / 'portal.css', media_type='text/css')

    @app.get('/portal.js')
    def js():
        return FileResponse(directory / 'portal.js', media_type='application/javascript')

    @app.get('/api/estudio')
    def orders(user=Depends(admin)):
        data = read_production()
        fields = [normalized(h) for h in data['headers']]
        with database() as db:
            states = {r['source_row']: dict(r) for r in db.execute('SELECT source_row,version,decision,published FROM team_workspaces')}
            requests = [dict(r) for r in db.execute('SELECT * FROM team_requests ORDER BY id DESC LIMIT 100')]
        rows = []
        for row in data['rows']:
            values = dict(zip(fields, row['values']))
            if values.get('ORDEN') and values.get('NOMBRE DEL CLIENTE'):
                rows.append(dict(row=row['source_row'], order=values['ORDEN'], client=values['NOMBRE DEL CLIENTE'], reference=values.get('REFERENCIA', ''), **{k:v for k,v in states.get(row['source_row'], {}).items() if k != 'source_row'}))
        return dict(orders=rows, requests=requests)

    @app.get('/api/estudio/{source_row}')
    def get_order(source_row: int, user=Depends(admin)):
        return details(source_row)

    @app.post('/api/estudio/{source_row}/sincronizar')
    def sync_designs(source_row: int, user=Depends(admin)):
        info = row_info(source_row)
        files, *_ = row_files(source_row, excel_only=True)
        images, message = excel_designs(source_row, files)
        if not images:
            raise HTTPException(422, message or 'El listado Excel no tiene diseños disponibles')
        designs = []
        with database() as db:
            db.execute('BEGIN IMMEDIATE')
            record = workspace(db, source_row, info)
            for number, mime, content in images[:4]:
                digest = hashlib.sha256(content).hexdigest()
                if mime not in ('image/jpeg', 'image/png', 'image/webp') or len(content) > 15_000_000:
                    raise HTTPException(422, 'Formato de imagen no admitido o imagen mayor de 15 MB')
                db.execute('INSERT OR IGNORE INTO team_images VALUES(?,?,?)', (digest, mime, content))
                designs.append(dict(number=number, digest=digest))
            serialized = json.dumps(designs, sort_keys=True)
            if serialized != record['designs']:
                db.execute("UPDATE team_workspaces SET designs=?,version=version+1,decision='pending' WHERE source_row=?", (serialized, source_row))
                event(db, source_row, 'Nueva versión de diseño', user)
        return details(source_row)

    @app.post('/api/estudio/{source_row}/enlace')
    def share(source_row: int, user=Depends(admin)):
        info, token = row_info(source_row), secrets.token_urlsafe(32)
        with database() as db:
            workspace(db, source_row, info)
            db.execute('DELETE FROM team_links WHERE source_row=?', (source_row,))
            db.execute('INSERT INTO team_links VALUES(?,?,?)', (hashlib.sha256(token.encode()).hexdigest(), source_row, time.time() + 30 * 86400))
            event(db, source_row, 'Enlace privado renovado · 30 días', user)
        return dict(url=f'/equipo/{token}')

    @app.delete('/api/estudio/{source_row}/enlace')
    def revoke(source_row: int, user=Depends(admin)):
        with database() as db:
            db.execute('DELETE FROM team_links WHERE source_row=?', (source_row,))
            event(db, source_row, 'Enlaces revocados', user)
        return dict(ok=True)

    @app.post('/api/estudio/{source_row}/catalogo')
    def publish(source_row: int, payload: dict = Body(...), user=Depends(admin)):
        info = row_info(source_row)
        sport = str(payload.get('sport', 'Fútbol'))
        if sport not in ('Fútbol', 'Baloncesto', 'Voleibol', 'Ciclismo', 'Running', 'Otros') or not isinstance(payload.get('published'), bool):
            raise HTTPException(422, 'Selecciona deporte y visibilidad')
        with database() as db:
            record = workspace(db, source_row, info)
            if payload['published'] and not json.loads(record['designs']):
                raise HTTPException(422, 'Primero importa un diseño del Excel')
            db.execute('UPDATE team_workspaces SET published=?,sport=? WHERE source_row=?', (int(payload['published']), sport, source_row))
            event(db, source_row, 'Visibilidad del catálogo', user, 'Público' if payload['published'] else 'Privado')
        return dict(ok=True)

    def image_response(source_row, digest, public=False):
        info = row_info(source_row)
        with database() as db:
            record = workspace(db, source_row, info)
            if (public and not record['published']) or digest not in [d['digest'] for d in json.loads(record['designs'])]:
                raise HTTPException(404)
            image = db.execute('SELECT mime,content FROM team_images WHERE digest=?', (digest,)).fetchone()
        if not image:
            raise HTTPException(404)
        return Response(image['content'], media_type=image['mime'], headers={'Cache-Control': 'private, no-store', 'X-Content-Type-Options': 'nosniff'})

    @app.get('/api/estudio/{source_row}/imagen/{digest}')
    def admin_image(source_row: int, digest: str, user=Depends(admin)):
        return image_response(source_row, digest)

    @app.get('/api/equipo/{token}/imagen/{digest}')
    def client_image(token: str, digest: str):
        return image_response(token_row(token), digest)

    @app.get('/api/equipo/{token}')
    def client_order(token: str):
        return details(token_row(token), token)

    @app.post('/api/equipo/{token}/decision')
    def decision(token: str, payload: dict = Body(...)):
        source_row = token_row(token)
        info = row_info(source_row)
        actor, note, choice = (str(payload.get(k, '')).strip() for k in ('actor', 'note', 'choice'))
        if not actor or len(actor) > 80 or len(note) > 2000 or choice not in ('approved', 'changes') or (choice == 'changes' and not note):
            raise HTTPException(422, 'Indica tu nombre y el motivo si solicitas cambios')
        with database() as db:
            db.execute('BEGIN IMMEDIATE')
            record = workspace(db, source_row, info)
            if not record['version'] or payload.get('version') != record['version']:
                raise HTTPException(409, 'El diseño cambió. Actualiza la página antes de aprobar.')
            db.execute('UPDATE team_workspaces SET decision=? WHERE source_row=?', (choice, source_row))
            event(db, source_row, f"{'Aprobado' if choice == 'approved' else 'Cambios solicitados'} · V{record['version']}", actor, note)
        return dict(ok=True)

    @app.put('/api/equipo/{token}/jugadores')
    def players(token: str, payload: dict = Body(...)):
        source_row = token_row(token)
        info = row_info(source_row)
        values = validate_players(payload.get('players'))
        with database() as db:
            db.execute('BEGIN IMMEDIATE')
            record = workspace(db, source_row, info)
            if payload.get('revision') != record['revision']:
                raise HTTPException(409, 'Otra persona actualizó el listado. Recarga antes de guardar.')
            db.execute('UPDATE team_workspaces SET players=?,revision=revision+1 WHERE source_row=?', (json.dumps(values, ensure_ascii=False), source_row))
            event(db, source_row, 'Listado de jugadores guardado', 'Enlace de cliente', f'{len(values)} jugadores')
        return dict(ok=True)

    @app.post('/api/equipo/{token}/importar')
    async def import_players(token: str, file: UploadFile = File(...)):
        source_row = token_row(token)
        row_info(source_row)
        content = await file.read(2_000_001)
        if len(content) > 2_000_000 or not (file.filename or '').lower().endswith('.xlsx'):
            raise HTTPException(422, 'Usa un archivo .xlsx de máximo 2 MB')
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20_000_000:
                    raise ValueError('Archivo demasiado grande')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                rows = list(book.active.iter_rows(max_row=502, max_col=3, values_only=True))
                if [normalized(v) for v in rows[0]] != ['NOMBRE', 'DORSAL', 'TALLA']:
                    raise ValueError('La primera fila debe ser NOMBRE, DORSAL, TALLA')
                values = [dict(name=r[0] or '', number=str(r[1]) if r[1] is not None else '', size=r[2] or '') for r in rows[1:] if any(v is not None for v in r)]
                return dict(players=validate_players(values))
            finally:
                book.close()
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, 'No se pudo leer el Excel. Revisa las columnas NOMBRE, DORSAL, TALLA.') from exc

    @app.post('/api/equipo/{token}/repetir')
    def reorder(token: str, payload: dict = Body(...)):
        source_row = token_row(token)
        row_info(source_row)
        quantity = payload.get('quantity')
        note = str(payload.get('note', '')).strip()
        if type(quantity) is not int or not 1 <= quantity <= 10000 or len(note) > 2000:
            raise HTTPException(422, 'Cantidad válida: 1–10000')
        with database() as db:
            if db.execute('SELECT 1 FROM team_requests WHERE source_row=? AND created>?', (source_row, time.time()-60)).fetchone():
                raise HTTPException(429, 'La solicitud ya fue recibida. Espera un minuto.')
            db.execute('INSERT INTO team_requests(source_row,kind,detail,created) VALUES(?,?,?,?)', (source_row, 'Repetición', json.dumps(dict(quantity=quantity, note=note)), time.time()))
        return dict(ok=True)

    @app.get('/api/catalogo')
    def catalog():
        with database() as db:
            records = [dict(r) for r in db.execute('SELECT * FROM team_workspaces WHERE published=1')]
        result = []
        for record in records:
            try:
                info = row_info(record['source_row'])
                if record['identity'] != identity(info):
                    continue
                designs = json.loads(record['designs'])
                result.append(dict(row=record['source_row'], title=info['reference'], sport=record['sport'], images=[dict(number=d['number'], url=f"/api/catalogo/{record['source_row']}/imagen/{d['digest']}") for d in designs]))
            except HTTPException:
                continue
        return result

    @app.get('/api/catalogo/{source_row}/imagen/{digest}')
    def catalog_image(source_row: int, digest: str):
        return image_response(source_row, digest, public=True)
