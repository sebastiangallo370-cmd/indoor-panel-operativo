"""La orden de REPROGRAMACIONES se verifica contra el PDF: unidades del Excel = Total unidades."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from app import verificar_remision as vr


def libro(ruta: Path, hojas: dict[str, int]):
    wb = Workbook()
    wb.active.title = 'REFERENCIA'
    for nombre, n in hojas.items():
        ws = wb.create_sheet(nombre)
        for i in range(n):
            ws.cell(row=8 + i, column=4, value='M')
    wb.save(ruta)


class Verificar(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.xlsx = Path(self.dir.name) / 'o.xlsx'

    def tearDown(self):
        self.dir.cleanup()

    def test_total_del_pdf(self):
        self.assertEqual(vr.total_del_pdf('algo\nTotal unidades: 1.245\nmás'), 1245)
        self.assertEqual(vr.total_del_pdf('TOTAL UNIDADES 30'), 30)
        self.assertIsNone(vr.total_del_pdf('sin total'))

    def test_cuadra(self):
        libro(self.xlsx, {'A200FUT': 20, 'A300CA01': 10})
        r = vr.verificar('Total unidades: 30', [{'cantidad': 20}, {'cantidad': 10}], self.xlsx)
        self.assertTrue(r['ok'])
        self.assertEqual((r['excel'], r['pdf']), (30, 30))
        self.assertIn('Verificado', r['mensaje'])

    def test_no_cuadra_si_falta_una_fila(self):
        libro(self.xlsx, {'A200FUT': 19, 'A300CA01': 10})
        r = vr.verificar('Total unidades: 30', [{'cantidad': 20}, {'cantidad': 10}], self.xlsx)
        self.assertFalse(r['ok'])
        self.assertIn('29', r['mensaje'])

    def test_item_que_no_se_leyo(self):
        libro(self.xlsx, {'A200FUT': 20})
        r = vr.verificar('Total unidades: 30', [{'cantidad': 20}], self.xlsx)
        self.assertFalse(r['ok'])
        self.assertIn('puede faltar algún ítem', r['mensaje'])

    def test_sin_total_usa_la_suma_de_items(self):
        libro(self.xlsx, {'A200FUT': 20})
        self.assertTrue(vr.verificar('', [{'cantidad': 20}], self.xlsx)['ok'])


if __name__ == '__main__':
    unittest.main()
